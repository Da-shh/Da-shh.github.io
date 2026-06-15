# -*- coding: utf-8 -*-
"""
Применяет к ГОТОВОМУ файлу презентации три правки, не трогая остальное
(в т.ч. ваши изменения размера шрифта):

  1) убирает подпись в левом верхнем углу («ИС «Тропа» · защита ВКР» /
     «Бакалаврская работа»);
  2) вставляет в левый верхний угол эмблему университета (если указан файл);
  3) увеличивает подписи в правом верхнем углу (раздел) и в правом нижнем
     углу (нумерация «NN / 17»).

Запуск:
    python3 apply_branding.py ВХОД.pptx ВЫХОД.pptx [logo.png]

Пример:
    python3 apply_branding.py "Презентация_ВКР_Тропа.pptx" "Готово.pptx" assets/tsu_logo.png
"""

import os
import re
import sys

from pptx import Presentation
from pptx.util import Inches, Pt

# Настройки
RIGHT_TOP_SIZE = 15.0     # размер подписи-раздела (правый верх)
FOOT_SIZE = 13.5          # размер нумерации (правый низ)
LOGO_LEFT_IN = 0.55
LOGO_TOP_IN = 0.30
LOGO_HEIGHT_IN = 0.62

FOOT_RE = re.compile(r"^\s*\d+\s*/\s*\d+\s*$")


def norm(text):
    return " ".join(text.upper().split())


def is_left_kicker(text):
    t = norm(text)
    return t.startswith("ИС «ТРОПА»") or t == "БАКАЛАВРСКАЯ РАБОТА" \
        or "ТРОПА» · ЗАЩИТА" in t


def set_size(shape, size_pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size_pt)
        # абзац без явных ранов
        if not p.runs and p.text:
            p.font.size = Pt(size_pt)


def emu_to_in(v):
    return (v or 0) / 914400.0


def process(in_path, out_path, logo_path=None):
    prs = Presentation(in_path)
    sw = emu_to_in(prs.slide_width)
    removed = enlarged_top = enlarged_foot = logos = 0

    for slide in prs.slides:
        to_remove = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text or ""
            if not txt.strip():
                continue
            left_in = emu_to_in(sh.left)
            top_in = emu_to_in(sh.top)

            # 1) левая подпись — на удаление
            if is_left_kicker(txt) and top_in < 1.3 and left_in < 5.0:
                to_remove.append(sh)
                continue

            # 3a) правый верхний угол (раздел) — увеличить
            if top_in < 1.0 and left_in > sw / 2.0 - 0.5:
                set_size(sh, RIGHT_TOP_SIZE)
                enlarged_top += 1
                continue

            # 3b) правый нижний угол (нумерация) — увеличить
            if FOOT_RE.match(txt) and top_in > 5.5:
                set_size(sh, FOOT_SIZE)
                enlarged_foot += 1
                continue

        for sh in to_remove:
            sh._element.getparent().remove(sh._element)
            removed += 1

        # 2) логотип в левый верхний угол
        if logo_path and os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(LOGO_LEFT_IN),
                                     Inches(LOGO_TOP_IN),
                                     height=Inches(LOGO_HEIGHT_IN))
            logos += 1

    prs.save(out_path)
    print(f"Готово: {out_path}")
    print(f"  убрано левых подписей : {removed}")
    print(f"  увеличено (правый верх): {enlarged_top}")
    print(f"  увеличено (правый низ) : {enlarged_foot}")
    print(f"  добавлено логотипов    : {logos}")
    if not logo_path:
        print("  (логотип не указан — добавьте файл и запустите снова, "
              "либо вставьте вручную)")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    inp, outp = sys.argv[1], sys.argv[2]
    logo = sys.argv[3] if len(sys.argv) > 3 else None
    process(inp, outp, logo)
